from pathlib import Path
import json,re
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
P=Path(__file__).resolve().parent; O=P.parents[2]/'outputs/nanda-demo-v5';O.mkdir(parents=True,exist_ok=True)
data=json.loads((P/'scenes.json').read_text());prompts={p['id']:p for p in json.loads((P/'prompts.json').read_text())['prompts']}
d=Document(); sec=d.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11);sec.top_margin=sec.bottom_margin=Inches(.7);sec.left_margin=sec.right_margin=Inches(.8)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
 s=d.styles[name];s.font.name='Arial';s.font.color.rgb=RGBColor(0,0,0)
d.styles['Normal'].font.size=Pt(11);d.styles['Normal'].paragraph_format.space_after=Pt(6);d.styles['Normal'].paragraph_format.line_spacing=1.08
for name in ['Heading 1','Heading 2']:
 d.styles[name].font.size=Pt(14 if name=='Heading 1' else 12);d.styles[name].paragraph_format.keep_with_next=True
h=d.add_paragraph('Bakery demo script',style='Title');h.paragraph_format.space_after=Pt(6)
d.add_paragraph('Revision 5 | Known agent ID resolution | October 19, 2026')
d.add_paragraph('Please comment on the numbered narration lines. Screen notes describe the animation and are not spoken. All screens, addresses and example codes are illustrative.')
d.add_paragraph('Alice runs the bakery in Muse. Bob is her customer in Wajo. The central example is a new message from Alice to Bob after his order, with his permission.')
line=1
for i,s in enumerate(data['scenes']):
 d.add_paragraph(f"{i+1}. {re.sub(r'^\d+\. ', '', s['title']).replace('’',chr(39))}",style='Heading 1')
 # Stable line numbers follow sentences, preserving the spoken p.m. abbreviation.
 parts=re.split(r"(?<=[.!?])\s+(?=[A-Z])|(?<=[.!?]')\s+(?=[A-Z])",s['narration'].replace("'Thanks. I'll","'Thanks§ I'll").replace('p.m.','p§m§'))
 for text in parts:
  d.add_paragraph(f'L{line:02d}  {text.replace(chr(167), chr(46))}');line+=1
 p=d.add_paragraph();p.add_run('On screen: ').bold=True;p.add_run(s['screen'])
 for pid in s['prompts']:
  p=d.add_paragraph();p.add_run(prompts[pid]['app']+' prompt: ').bold=True;p.add_run(prompts[pid]['text'])
# Sources and protocol detail are deliberately outside the spoken script.
d.add_paragraph('Technical notes for reviewers',style='Heading 1')
notes=[
'The walkthrough starts with known identities. It contains no capability search, ranked results or manual search step. @bakery.com is a user-facing shorthand for the known business domain.',
'The bakery is assumed to publish a domain-authorized contact path. Verifying alice@bakery.com proves control of that mailbox, not control of the entire domain.',
'For Bob, the lookup returns an authority-bound pointer to a hosted Agent Card. The card supplies the endpoint and authentication requirements. Bob proves control of his mailbox without controlling gmail.com.',
'Resolution and authorization are separate. Bob has opted into order updates; the sending agent must still follow current access rules. The later bakery message begins a new connection rather than only replying within an existing session.',
'Poppy and ARD use domain-anchored publications, but they are not identical protocols. Poppy also defines interaction and authorization behavior. This video borrows the simple consumer-to-business demonstration structure, then focuses on resolving a consumer identity for the reverse direction.',
'The narration says index and AgentBox. NANDA may appear on screen. Exact prompt text comes from the same JSON used by the animation and reading guide.'
]
for n in notes:d.add_paragraph(n)
def link(label,url):
 p=d.add_paragraph();r=p.add_run(label+': ');r.bold=True;p.add_run(url)
link('Resolution paper, revision 01','https://datatracker.ietf.org/doc/html/draft-raskar-agentic-web-federated-resolution-01')
link('Poppy specification, draft 0.2','https://personalagentprotocol.org/docs/0.2/spec')
link('Sierra explanation and demo','https://sierra.ai/blog/poppy')
link('Official Muse, Redfin and Rocket video','https://x.com/SierraPlatform/status/2108245197440307584')
f=sec.footer.paragraphs[0];f.alignment=2;r=f.add_run();field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');r._r.addnext(field)
d.core_properties.title='Bakery demo script, revision 5';d.core_properties.subject='Known agent ID resolution and business-to-consumer messaging';d.core_properties.author='Vedh Krishnan'
d.save(O/'bakery-demo-script-v5.docx');print(O/'bakery-demo-script-v5.docx')
